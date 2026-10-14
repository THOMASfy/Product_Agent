#!/usr/bin/env python3
"""统一文件生成器：根据 JSON 规格产出 PPTX / DOCX / XLSX / PDF。

用法:
    python build.py <spec.json>

规格 (spec.json) 公共字段:
{
  "type":   "pptx" | "docx" | "xlsx" | "pdf",
  "output": "data/user/<open_id>/output/文件名",
  "style":  { "primary": "0070C0", "secondary": "0445FE", "accent": "0FA4FE",
              "text": "333333", "title_font": "微软雅黑", "body_font": "微软雅黑",
              "title_size": 32, "body_size": 18 }
}

各类型专属字段见 .opencode/skills/docgen/SKILL.md。
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile

try:  # 保证中文路径/文本在任意控制台编码下都能打印。
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

# 仓库根目录 = <repo>/.opencode/skills/docgen/scripts/build.py -> 需要上溯 5 层
REPO_ROOT = os.path.abspath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "..")
)

DEFAULT_STYLE = {
    "primary": "0070C0",
    "secondary": "0445FE",
    "accent": "0FA4FE",
    "text": "333333",
    "title_font": "微软雅黑",
    "body_font": "微软雅黑",
    "title_size": 32,
    "body_size": 18,
}


def resolve(path):
    if not path:
        return path
    if os.path.isabs(path):
        return path
    return os.path.normpath(os.path.join(REPO_ROOT, path))


def merge_style(spec):
    style = dict(DEFAULT_STYLE)
    style.update(spec.get("style") or {})
    return style


# --------------------------------------------------------------------------- #
# 通用 / PPTX
# --------------------------------------------------------------------------- #
def _hex(value, fallback="000000"):
    value = (value or fallback).lstrip("#")
    return value.upper()


def build_pptx(spec, style, output):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    r_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

    def color(value, fallback="000000"):
        try:
            return RGBColor.from_string(_hex(value, fallback))
        except Exception:
            return RGBColor.from_string(fallback)

    def delete_all_slides(prs):
        sld_id_lst = prs.slides._sldIdLst
        for sld_id in list(sld_id_lst):
            try:
                prs.part.drop_rel(sld_id.get(r_id))
            except Exception:
                pass
            sld_id_lst.remove(sld_id)

    def pick_layout(prs, hint):
        layouts = list(prs.slide_layouts)
        if not layouts:
            return None
        keywords = {
            "title": ["标题幻灯片", "封面", "title"],
            "section": ["节标题", "分节", "section", "目录"],
            "content": ["标题和内容", "内容", "正文", "content"],
            "blank": ["空白", "blank"],
        }.get(hint, [])
        for kw in keywords:
            for layout in layouts:
                if kw.lower() in (layout.name or "").lower():
                    return layout
        fallback = {"title": 0, "content": 1, "section": 2, "blank": 6}.get(hint, 0)
        return layouts[fallback] if fallback < len(layouts) else layouts[0]

    def style_run(run, font, size, col, bold=False):
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color(col)

    def set_title(slide, text):
        if not text:
            return
        title = slide.shapes.title
        if title is not None:
            title.text = text
            for para in title.text_frame.paragraphs:
                for run in para.runs:
                    style_run(run, style["title_font"], style["title_size"], style["primary"], True)
            return
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.1))
        box.text_frame.text = text
        for para in box.text_frame.paragraphs:
            for run in para.runs:
                style_run(run, style["title_font"], style["title_size"], style["primary"], True)

    def first_body(slide):
        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.idx != 0 and placeholder.has_text_frame:
                return placeholder
        return None

    def set_bullets(slide, bullets):
        if not bullets:
            return
        body = first_body(slide) or slide.shapes.add_textbox(
            Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0)
        )
        frame = body.text_frame
        frame.word_wrap = True
        frame.text = bullets[0]
        for extra in bullets[1:]:
            frame.add_paragraph().text = extra
        for para in frame.paragraphs:
            para.alignment = PP_ALIGN.LEFT
            for run in para.runs:
                style_run(run, style["body_font"], style["body_size"], style["text"])

    def add_subtitle(slide, text):
        if not text:
            return
        holder = first_body(slide)
        if holder is not None:
            holder.text = text
            for para in holder.text_frame.paragraphs:
                for run in para.runs:
                    style_run(run, style["body_font"], style["body_size"], style["secondary"])
            return
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.2))
        box.text_frame.text = text
        for para in box.text_frame.paragraphs:
            for run in para.runs:
                style_run(run, style["body_font"], style["body_size"], style["secondary"])

    template = resolve(spec.get("template"))
    if template and os.path.exists(template):
        prs = Presentation(template)
        delete_all_slides(prs)
        print(f"using template: {template} ({len(prs.slide_layouts)} layouts)")
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        if template:
            print(f"template not found, building blank: {template}", file=sys.stderr)

    for item in spec.get("slides", []):
        layout = pick_layout(prs, item.get("layout", "content"))
        slide = prs.slides.add_slide(layout)
        set_title(slide, item.get("title"))
        if item.get("layout") == "title" and item.get("subtitle"):
            add_subtitle(slide, item["subtitle"])
        set_bullets(slide, item.get("bullets"))
        if item.get("notes"):
            slide.notes_slide.notes_text_frame.text = item["notes"]

    prs.save(output)


# --------------------------------------------------------------------------- #
# DOCX / PDF
# --------------------------------------------------------------------------- #
def _set_cjk(run, font):
    from docx.oxml.ns import qn

    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def build_docx(spec, style, output):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = style["body_font"]
    normal.font.size = Pt(style["body_size"])
    normal.font.color.rgb = __import__("docx").shared.RGBColor.from_string(_hex(style["text"]))

    title = spec.get("title")
    if title:
        heading = doc.add_heading(title, level=0)
        for run in heading.runs:
            _set_cjk(run, style["title_font"])
            run.font.size = Pt(max(style["title_size"], 28))
            run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(_hex(style["primary"]))
    if spec.get("subtitle"):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(spec["subtitle"])
        _set_cjk(run, style["body_font"])
        run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(_hex(style["secondary"]))

    for section in spec.get("sections", []):
        if section.get("heading"):
            heading = doc.add_heading(section["heading"], level=int(section.get("level", 1)))
            for run in heading.runs:
                _set_cjk(run, style["title_font"])
                run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(_hex(style["primary"]))
        for text in section.get("paragraphs", []):
            para = doc.add_paragraph()
            run = para.add_run(text)
            _set_cjk(run, style["body_font"])
        for text in section.get("bullets", []):
            para = doc.add_paragraph(style="List Bullet")
            run = para.add_run(text)
            _set_cjk(run, style["body_font"])

    # 顶层便捷写法：直接给 paragraphs / bullets。
    for text in spec.get("paragraphs", []):
        run = doc.add_paragraph().add_run(text)
        _set_cjk(run, style["body_font"])
    for text in spec.get("bullets", []):
        run = doc.add_paragraph(style="List Bullet").add_run(text)
        _set_cjk(run, style["body_font"])

    doc.save(output)


def build_pdf(spec, style, output):
    """先构建 DOCX，再用 LibreOffice 转 PDF（对中文最稳妥）。"""
    docx_path = os.path.join(tempfile.mkdtemp(prefix="docgen_"), "doc.docx")
    build_docx(spec, style, docx_path)
    soffice = find_soffice()
    if not soffice:
        raise SystemExit("生成 PDF 需要 LibreOffice (soffice)，请先安装")
    out_dir = os.path.dirname(output) or "."
    subprocess.run(
        [soffice, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
        check=True,
        timeout=180,
    )
    produced = os.path.join(out_dir, "doc.pdf")
    if not os.path.exists(produced):
        raise SystemExit("PDF 转换失败")
    if os.path.abspath(produced) != os.path.abspath(output):
        shutil.move(produced, output)
    shutil.rmtree(os.path.dirname(docx_path), ignore_errors=True)


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #
def build_xlsx(spec, style, output):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor=_hex(style["primary"]))
    header_font = Font(name=style["body_font"], bold=True, color="FFFFFF", size=int(style["body_size"]))
    cell_font = Font(name=style["body_font"], color=_hex(style["text"]), size=int(style["body_size"]))

    sheets = spec.get("sheets") or [{"name": spec.get("title") or "Sheet1", "rows": spec.get("rows", [])}]
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(title=(sheet.get("name") or "Sheet")[:31])
        rows = sheet.get("rows", [])
        has_header = bool(sheet.get("header", True)) and bool(rows)
        for row in rows:
            ws.append(row)
        if has_header:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.freeze_panes = "A2"
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = cell_font
        for column_cells in ws.columns:
            width = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 10), 60)
    wb.save(output)


# --------------------------------------------------------------------------- #
# LibreOffice 辅助
# --------------------------------------------------------------------------- #
SOFFICE_CANDIDATES = [
    os.environ.get("SOFFICE_PATH", ""),
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "soffice",
]


def find_soffice():
    for candidate in SOFFICE_CANDIDATES:
        if not candidate:
            continue
        if os.sep in candidate or "/" in candidate:
            if os.path.exists(candidate):
                return candidate
        else:
            found = shutil.which(candidate)
            if found:
                return found
    return None


BUILDERS = {
    "pptx": build_pptx,
    "docx": build_docx,
    "xlsx": build_xlsx,
    "pdf": build_pdf,
}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("spec", help="path to spec json")
    args = parser.parse_args()

    with open(args.spec, "r", encoding="utf-8") as handle:
        spec = json.load(handle)

    kind = (spec.get("type") or "").lower()
    if kind not in BUILDERS:
        raise SystemExit(f"spec.type 必须是 {', '.join(BUILDERS)} 之一")

    output = resolve(spec.get("output"))
    if not output:
        raise SystemExit("spec.output is required")
    os.makedirs(os.path.dirname(os.path.abspath(output)), exist_ok=True)

    BUILDERS[kind](spec, merge_style(spec), output)
    print(f"SAVED {os.path.abspath(output)}")


if __name__ == "__main__":
    main()
